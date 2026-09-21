"""Add a video link line near the top of PI-Transfer-Tracker-Getting-Started.docx.
Idempotent: safe to run multiple times; only inserts if not already present.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from copy import deepcopy

DOC_PATH = '/home/abrown/transfer-tracker/PI-Transfer-Tracker-Getting-Started.docx'
VIDEO_URL = 'https://pharmacyinnovationsnet-my.sharepoint.com/:v:/g/personal/abrown_pharmacyinnovations_net/IQA6K8ip3Nb4RY_wDcr6xUseARyBqqddL3X1hHSt98geshE?e=YfQb55'
MARKER = 'Training video:'

doc = Document(DOC_PATH)

# Check if already added
for p in doc.paragraphs:
    if MARKER in p.text:
        print(f'Already present — no changes made.')
        raise SystemExit(0)

# Find the "Full help:" paragraph and insert after it
target_idx = None
for i, p in enumerate(doc.paragraphs):
    if 'Full help' in p.text or 'help.html' in p.text:
        target_idx = i
        break

if target_idx is None:
    # Fallback: insert after the Website line, or at the start
    for i, p in enumerate(doc.paragraphs):
        if 'Website:' in p.text or 'red-island' in p.text:
            target_idx = i
            break

if target_idx is None:
    print('Could not find insertion point (Website/Full help line). Aborting.')
    raise SystemExit(1)

# Insert a new paragraph AFTER the target
target = doc.paragraphs[target_idx]
new_p = target.insert_paragraph_before('')
# The insert_paragraph_before puts it BEFORE target; we need to swap — easier to do XML manipulation
# Actually simpler approach: use the underlying XML to insert after
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Remove the empty paragraph we just made
new_p._element.getparent().remove(new_p._element)

# Build a new paragraph and place it right after the target
new_para = OxmlElement('w:p')

# Bold "📹 Training video: " prefix
r1 = OxmlElement('w:r')
rpr1 = OxmlElement('w:rPr')
b = OxmlElement('w:b')
rpr1.append(b)
r1.append(rpr1)
t1 = OxmlElement('w:t')
t1.text = '📹 Training video: '
t1.set(qn('xml:space'), 'preserve')
r1.append(t1)
new_para.append(r1)

# Hyperlink to the video (as a plain-text URL run — simplest reliable approach)
r2 = OxmlElement('w:r')
rpr2 = OxmlElement('w:rPr')
color = OxmlElement('w:color')
color.set(qn('w:val'), '2A6EBB')
rpr2.append(color)
u = OxmlElement('w:u')
u.set(qn('w:val'), 'single')
rpr2.append(u)
r2.append(rpr2)
t2 = OxmlElement('w:t')
t2.text = 'Watch the PI Transfer Tracker training video'
r2.append(t2)
new_para.append(r2)

# Rest of the sentence
r3 = OxmlElement('w:r')
t3 = OxmlElement('w:t')
t3.text = f' ({VIDEO_URL}) — ~10 min screen-recorded walkthrough of the most common tasks. Watch this first before your day-one setup.'
t3.set(qn('xml:space'), 'preserve')
r3.append(t3)
new_para.append(r3)

# Insert after the target paragraph
target._element.addnext(new_para)

doc.save(DOC_PATH)
print(f'Added video link line after paragraph {target_idx}: "{target.text[:60]}"')
