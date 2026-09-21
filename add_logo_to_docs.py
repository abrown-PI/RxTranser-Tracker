"""Add PI landscape logo to the top of the Transfer Tracker docs.
Idempotent: skips if the logo is already present.
"""
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

LOGO_PATH = '/mnt/c/Users/ABrown/OneDrive - Pharmacy Innovations/Documents/Tina/Flyers/Pharmacy Innovations landscape logo.png'
DOCS = [
    '/home/abrown/transfer-tracker/PI-Transfer-Tracker-Getting-Started.docx',
    '/mnt/c/Users/ABrown/OneDrive - Pharmacy Innovations/Documents/Portals/Transfer Tracker/Transfer-Tracker-Go-Live-Notice.docx',
]

def has_image(doc):
    """Check if the doc already has an inline image in the first few paragraphs."""
    for p in doc.paragraphs[:5]:
        if p._element.findall('.//' + qn('w:drawing')):
            return True
    return False

for doc_path in DOCS:
    if not os.path.exists(doc_path):
        print(f'MISSING: {doc_path}')
        continue

    doc = Document(doc_path)

    if has_image(doc):
        print(f'SKIP (logo already present): {os.path.basename(doc_path)}')
        continue

    # Insert a new paragraph at the very beginning, add the logo image, centered
    body = doc.element.body
    first_p = doc.paragraphs[0]

    # Create a new paragraph BEFORE the first one for the logo
    logo_para = first_p.insert_paragraph_before('')
    logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_run = logo_para.add_run()
    logo_run.add_picture(LOGO_PATH, width=Inches(2.5))

    doc.save(doc_path)
    print(f'ADDED logo to: {os.path.basename(doc_path)}')
